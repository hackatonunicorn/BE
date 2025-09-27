"""
Модуль для анализа питч-деков с использованием Claude API
"""
import os
import io
import json
import asyncio
from typing import Dict, Any, Optional, List, Union
from pathlib import Path
import logging
from datetime import datetime

# Импорты для обработки файлов
try:
    import PyPDF2
    from pptx import Presentation
    from docx import Document
except ImportError:
    PyPDF2 = None
    Presentation = None
    Document = None

# Импорты для AI
try:
    import anthropic
except ImportError:
    anthropic = None

from app.ai.text_processing import text_processor
from app.core.config import settings

logger = logging.getLogger(__name__)


class FileExtractor:
    """Класс для извлечения текста из различных типов файлов"""
    
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Извлечение текста из PDF файла"""
        if not PyPDF2:
            raise ImportError("PyPDF2 not installed. Install with: pip install PyPDF2")
        
        try:
            text = ""
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    text += page.extract_text() + "\n"
            
            return text
        except Exception as e:
            logger.error(f"Error extracting text from PDF {file_path}: {e}")
            raise

    @staticmethod
    def extract_text_from_pptx(file_path: str) -> str:
        """Извлечение текста из PowerPoint файла"""
        if not Presentation:
            raise ImportError("python-pptx not installed. Install with: pip install python-pptx")
        
        try:
            text = ""
            prs = Presentation(file_path)
            
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
            
            return text
        except Exception as e:
            logger.error(f"Error extracting text from PPTX {file_path}: {e}")
            raise

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Извлечение текста из Word документа"""
        if not Document:
            raise ImportError("python-docx not installed. Install with: pip install python-docx")
        
        try:
            doc = Document(file_path)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            return text
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_path}: {e}")
            raise

    @staticmethod
    def extract_text_from_file(file_path: str) -> str:
        """Универсальное извлечение текста из файла по расширению"""
        file_extension = Path(file_path).suffix.lower()
        
        if file_extension == '.pdf':
            return FileExtractor.extract_text_from_pdf(file_path)
        elif file_extension == '.pptx':
            return FileExtractor.extract_text_from_pptx(file_path)
        elif file_extension == '.docx':
            return FileExtractor.extract_text_from_docx(file_path)
        elif file_extension == '.txt':
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")


class PitchAnalyzer:
    """Класс для анализа питч-деков с использованием Claude API"""
    
    def __init__(self, claude_api_key: Optional[str] = None):
        self.claude_api_key = claude_api_key or os.getenv('CLAUDE_API_KEY') or getattr(settings, 'CLAUDE_API_KEY', None)
        
        if not self.claude_api_key and anthropic:
            logger.warning("Claude API key not provided. AI analysis will be simulated.")
        
        self.client = None
        if self.claude_api_key and anthropic:
            try:
                self.client = anthropic.Anthropic(api_key=self.claude_api_key)
            except Exception as e:
                logger.error(f"Failed to initialize Claude client: {e}")
        
        self.file_extractor = FileExtractor()

    def extract_text_from_pitch_deck(self, file_path: str) -> str:
        """Извлечение текста из питч-дека"""
        try:
            if not os.path.exists(file_path):
                raise FileNotFoundError(f"File not found: {file_path}")
            
            text = self.file_extractor.extract_text_from_file(file_path)
            
            # Предварительная обработка текста
            processed_text = text_processor.preprocess_for_ai_analysis(text)
            
            logger.info(f"Successfully extracted {len(processed_text)} characters from {file_path}")
            return processed_text
            
        except Exception as e:
            logger.error(f"Error extracting text from pitch deck {file_path}: {e}")
            raise

    def _create_analysis_prompt(self, text: str) -> str:
        """Создание промпта для анализа питч-дека"""
        return f"""
Analyze the following pitch deck content and extract key information. Please provide a structured analysis in JSON format.

PITCH DECK CONTENT:
{text}

Please analyze and extract the following information in valid JSON format:

{{
    "industry": "string - main industry/sector (e.g., 'Technology', 'FinTech', 'HealthTech', 'CleanTech')",
    "stage": "string - development stage (e.g., 'Seed', 'Series A', 'Series B', 'Series C')",
    "funding_amount": "integer - requested funding amount in USD (e.g., 2000000 for $2M)",
    "market_size": "string - total addressable market size (e.g., '$50B', 'Large', 'Growing')",
    "key_metrics": {{
        "revenue": "number - current or projected revenue",
        "users": "number - number of users/customers",
        "growth_rate": "number - growth rate percentage",
        "other_metrics": "object - any other important metrics"
    }},
    "team_info": {{
        "team_size": "number - number of team members",
        "founders": "array - list of founder names/roles",
        "key_backgrounds": "array - relevant backgrounds/experience"
    }},
    "business_model": "string - description of how the company makes money",
    "competitive_advantages": "array - list of key competitive advantages or unique selling points",
    "problem_solution": {{
        "problem": "string - problem being solved",
        "solution": "string - proposed solution"
    }},
    "financial_projections": {{
        "revenue_projection": "string - revenue projections if mentioned",
        "break_even": "string - break-even timeline if mentioned"
    }},
    "use_of_funds": "string - how the raised funds will be used",
    "confidence_score": "number - confidence score from 0.0 to 1.0 for the analysis quality"
}}

Focus on extracting factual information. If some information is not available in the text, use null for that field. Be precise with numbers and avoid speculation.
"""

    async def analyze_with_claude(self, text: str) -> Dict[str, Any]:
        """Анализ текста с использованием Claude API"""
        if not self.client:
            logger.warning("Claude client not available, using rule-based analysis")
            return self._rule_based_analysis(text)
        
        try:
            prompt = self._create_analysis_prompt(text)
            
            response = await asyncio.create_task(
                asyncio.to_thread(
                    self.client.messages.create,
                    model="claude-3-sonnet-20240229",
                    max_tokens=2000,
                    messages=[{
                        "role": "user",
                        "content": prompt
                    }]
                )
            )
            
            # Извлекаем контент ответа
            content = response.content[0].text if response.content else ""
            
            # Пытаемся парсить JSON из ответа
            try:
                # Ищем JSON в ответе (может быть обернут в ```json```)
                json_match = content
                if "```json" in content:
                    json_start = content.find("```json") + 7
                    json_end = content.find("```", json_start)
                    json_match = content[json_start:json_end].strip()
                elif "{" in content and "}" in content:
                    json_start = content.find("{")
                    json_end = content.rfind("}") + 1
                    json_match = content[json_start:json_end]
                
                result = json.loads(json_match)
                
                # Валидация и очистка результата
                result = self._validate_and_clean_result(result)
                result['analysis_method'] = 'claude_ai'
                result['analyzed_at'] = datetime.now().isoformat()
                
                logger.info("Successfully analyzed pitch deck with Claude API")
                return result
                
            except json.JSONDecodeError as e:
                logger.error(f"Failed to parse Claude response as JSON: {e}")
                logger.debug(f"Claude response content: {content}")
                return self._rule_based_analysis(text)
                
        except Exception as e:
            logger.error(f"Error analyzing with Claude API: {e}")
            return self._rule_based_analysis(text)

    def _rule_based_analysis(self, text: str) -> Dict[str, Any]:
        """Анализ на основе правил (fallback если Claude API недоступен)"""
        logger.info("Performing rule-based analysis")
        
        # Используем text_processor для извлечения информации
        industry_info = text_processor.detect_industry(text)
        stage_info = text_processor.detect_stage(text)
        funding_info = text_processor.extract_funding_amount(text)
        metrics = text_processor.extract_key_metrics(text)
        team_info = text_processor.extract_team_info(text)
        business_keywords = text_processor.extract_business_model_keywords(text)
        advantages = text_processor.extract_competitive_advantages(text)
        
        result = {
            'industry': industry_info.get('industry', 'unknown'),
            'stage': stage_info.get('stage', 'unknown'),
            'funding_amount': funding_info.get('amount') if funding_info else None,
            'market_size': metrics.get('market_size', {}).get('value') if 'market_size' in metrics else None,
            'key_metrics': {
                'revenue': metrics.get('revenue', {}).get('value') if 'revenue' in metrics else None,
                'users': metrics.get('users', {}).get('value') if 'users' in metrics else None,
                'growth_rate': metrics.get('growth', {}).get('value') if 'growth' in metrics else None,
                'other_metrics': {k: v for k, v in metrics.items() if k not in ['revenue', 'users', 'growth']}
            },
            'team_info': {
                'team_size': int(team_info.get('team_size', [0])[0]) if team_info.get('team_size') else None,
                'founders': team_info.get('founders', []),
                'key_backgrounds': team_info.get('background', [])
            },
            'business_model': ', '.join(business_keywords) if business_keywords else 'unknown',
            'competitive_advantages': advantages,
            'problem_solution': {
                'problem': None,  # Сложно извлечь правилами
                'solution': None
            },
            'financial_projections': {
                'revenue_projection': None,
                'break_even': None
            },
            'use_of_funds': None,
            'confidence_score': self._calculate_confidence_score(industry_info, stage_info, funding_info, metrics),
            'analysis_method': 'rule_based',
            'analyzed_at': datetime.now().isoformat()
        }
        
        return self._validate_and_clean_result(result)

    def _calculate_confidence_score(self, industry_info, stage_info, funding_info, metrics) -> float:
        """Расчет уровня уверенности в анализе"""
        score = 0.0
        
        # Очки за найденную информацию
        if industry_info.get('confidence', 0) > 0:
            score += 0.2
        if stage_info.get('confidence', 0) > 0:
            score += 0.2
        if funding_info:
            score += 0.3
        if metrics:
            score += 0.2 * min(len(metrics), 2)  # Максимум 0.4 за метрики
        
        return min(score, 1.0)

    def _validate_and_clean_result(self, result: Dict[str, Any]) -> Dict[str, Any]:
        """Валидация и очистка результата анализа"""
        # Определяем схему с типами по умолчанию
        schema = {
            'industry': str,
            'stage': str,
            'funding_amount': int,
            'market_size': str,
            'key_metrics': dict,
            'team_info': dict,
            'business_model': str,
            'competitive_advantages': list,
            'problem_solution': dict,
            'financial_projections': dict,
            'use_of_funds': str,
            'confidence_score': float
        }
        
        # Валидируем и приводим типы
        for key, expected_type in schema.items():
            if key not in result:
                result[key] = None
            elif result[key] is not None:
                try:
                    if expected_type == int and isinstance(result[key], (str, float)):
                        # Пытаемся конвертировать строки в числа
                        if isinstance(result[key], str):
                            # Удаляем символы валюты и запятые
                            clean_value = result[key].replace('$', '').replace(',', '').strip()
                            result[key] = int(float(clean_value))
                        else:
                            result[key] = int(result[key])
                    elif expected_type == float and isinstance(result[key], (str, int)):
                        result[key] = float(result[key])
                    elif expected_type == list and not isinstance(result[key], list):
                        result[key] = [result[key]] if result[key] else []
                    elif expected_type == dict and not isinstance(result[key], dict):
                        result[key] = {}
                except (ValueError, TypeError) as e:
                    logger.warning(f"Failed to convert {key} to {expected_type}: {e}")
                    result[key] = None
        
        # Убеждаемся, что словари имеют правильную структуру
        if not isinstance(result.get('key_metrics'), dict):
            result['key_metrics'] = {}
        
        if not isinstance(result.get('team_info'), dict):
            result['team_info'] = {}
        
        if not isinstance(result.get('problem_solution'), dict):
            result['problem_solution'] = {'problem': None, 'solution': None}
        
        if not isinstance(result.get('financial_projections'), dict):
            result['financial_projections'] = {'revenue_projection': None, 'break_even': None}
        
        # Ограничиваем confidence_score от 0 до 1
        if 'confidence_score' in result and result['confidence_score'] is not None:
            result['confidence_score'] = max(0.0, min(1.0, float(result['confidence_score'])))
        
        return result

    async def analyze_pitch_deck_file(self, file_path: str) -> Dict[str, Any]:
        """Полный анализ питч-дека из файла"""
        try:
            # Извлекаем текст
            text = self.extract_text_from_pitch_deck(file_path)
            
            if not text.strip():
                raise ValueError("No text content found in the file")
            
            # Анализируем с помощью AI
            result = await self.analyze_with_claude(text)
            
            # Добавляем метаданные
            result['source_file'] = file_path
            result['extracted_text_length'] = len(text)
            
            logger.info(f"Successfully analyzed pitch deck: {file_path}")
            return result
            
        except Exception as e:
            logger.error(f"Error analyzing pitch deck file {file_path}: {e}")
            raise

    def analyze_pitch_deck_text(self, text: str) -> Dict[str, Any]:
        """Анализ питч-дека из текста (синхронная версия)"""
        try:
            # Обрабатываем текст
            processed_text = text_processor.preprocess_for_ai_analysis(text)
            
            if not processed_text.strip():
                raise ValueError("No meaningful text content provided")
            
            # Выполняем rule-based анализ (синхронно)
            result = self._rule_based_analysis(processed_text)
            
            # Добавляем метаданные
            result['source_type'] = 'text'
            result['extracted_text_length'] = len(processed_text)
            
            logger.info("Successfully analyzed pitch deck text")
            return result
            
        except Exception as e:
            logger.error(f"Error analyzing pitch deck text: {e}")
            raise


# Создаем глобальный экземпляр для использования
pitch_analyzer = PitchAnalyzer()
